from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt

MD = Path('../SLIDES_NOTES.md')
PR = Path('../Presentation_Project.pptx')

prs = Presentation()

lines = MD.read_text(encoding='utf-8').splitlines()

slide = None
for line in lines:
    if line.startswith('# '):
        # new section as title slide
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.shapes.title
        title.text = line[2:].strip()
    elif line.startswith('## '):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        title = slide.shapes.title
        title.text = line[3:].strip()
    elif line.startswith('- '):
        if slide is None:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = 'Points'
        body = slide.shapes.placeholders[1].text_frame
        p = body.add_paragraph()
        p.text = line[2:].strip()
        p.level = 1
    elif line.strip() == '':
        continue
    else:
        # plain text -> add to body
        if slide is None:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = ''
        body = slide.shapes.placeholders[1].text_frame
        p = body.add_paragraph()
        p.text = line.strip()

prs.save(PR)
print('Saved', PR)
