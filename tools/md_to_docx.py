from pathlib import Path
from docx import Document

MD = Path('../REPORT_PROJECT_SUMMARY.md')
DOCX = Path('../REPORT_PROJECT_SUMMARY.docx')

doc = Document()

lines = MD.read_text(encoding='utf-8').splitlines()
for line in lines:
    if line.startswith('# '):
        doc.add_heading(line[2:].strip(), level=1)
    elif line.startswith('## '):
        doc.add_heading(line[3:].strip(), level=2)
    elif line.startswith('### '):
        doc.add_heading(line[4:].strip(), level=3)
    elif line.strip() == '---':
        doc.add_page_break()
    else:
        doc.add_paragraph(line)

doc.save(DOCX)
print('Saved', DOCX)
